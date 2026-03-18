from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


DOWNLOADS = Path.home() / "Downloads"
OUTPUT_PATH = Path("src/data/monsterEp01Data.json")
STORYBOARD_NAME_HINT = "湖村_漫画转分镜"
SCHEDULE_NAME_HINT = "湖村_每日详细排期"

STATUS_LABEL_TO_KEY = {
    "完成": "done",
    "修改中": "revision",
    "未开始": "pending",
    "": "pending",
}

TYPE_META = {
    "★ 关键帧": {"key": "keyframe", "short": "关键帧"},
    "◆ 场景图": {"key": "scene", "short": "场景图"},
    "○ 普通图": {"key": "normal", "short": "普通图"},
}

MODULE_DEFAULTS = {
    "scriptDev": {
        "title": "剧本开发",
        "description": "主线梳理 / 分集大纲 / 分集剧本",
        "subitems": [
            {"key": "storyMainline", "label": "主线梳理", "status": "done"},
            {"key": "episodeOutline", "label": "分集大纲", "status": "done"},
            {"key": "episodeScript", "label": "分集剧本", "status": "active"},
        ],
    },
    "assetDev": {
        "title": "资产管理",
        "description": "人物设计 / 场景设计",
        "subitems": [
            {"key": "characterDesign", "label": "人物设计", "status": "active"},
            {"key": "sceneDesign", "label": "场景设计", "status": "active"},
        ],
    },
    "storyboardDev": {
        "title": "分镜制作",
        "description": "文字分镜补充 / AI 图片分镜",
        "subitems": [
            {"key": "episodeStoryboard", "label": "文字分镜补充", "status": "pending"},
            {"key": "comicToImage", "label": "AI 图片分镜", "status": "active"},
        ],
    },
    "postDev": {
        "title": "后期",
        "description": "视频制作 / 音乐后期",
        "subitems": [
            {"key": "videoProduction", "label": "视频制作", "status": "pending"},
            {"key": "musicPost", "label": "音乐后期", "status": "pending"},
        ],
    },
}

DEFAULTS = {
    "members": ["LU", "敏", "铭"],
    "moduleState": {
        "scriptDev": {
            "status": "active",
            "subitems": {
                "storyMainline": "done",
                "episodeOutline": "done",
                "episodeScript": "active",
            },
        },
        "assetDev": {
            "status": "done",
            "subitems": {
                "characterDesign": "active",
                "sceneDesign": "active",
            },
        },
        "storyboardDev": {
            "status": "active",
            "subitems": {
                "episodeStoryboard": "pending",
                "comicToImage": "active",
            },
        },
        "postDev": {
            "status": "pending",
            "subitems": {
                "videoProduction": "pending",
                "musicPost": "pending",
            },
        },
    },
}


def find_latest_file(name_hint: str) -> Path:
    candidates = [
        path
        for path in DOWNLOADS.glob("*.xlsx")
        if name_hint in path.name and not path.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(f"Could not find workbook containing {name_hint!r} in {DOWNLOADS}")
    return sorted(candidates, key=lambda item: item.stat().st_mtime, reverse=True)[0]


def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_mmdd(value: str) -> date:
    month, day = map(int, value.split("/"))
    return date(2026, month, day)


def parse_yyyy_mm_dd(value: str) -> date:
    year, month, day = map(int, value.split("/"))
    return date(year, month, day)


def iso_date(value: date) -> str:
    return value.isoformat()


def chinese_weekday(value: date) -> str:
    return f"周{'一二三四五六日'[value.weekday()]}"


def format_cn_date(value: date, include_weekday: bool = True) -> str:
    if include_weekday:
        return f"{value.year}年{value.month}月{value.day}日（{chinese_weekday(value)}）"
    return f"{value.year}年{value.month}月{value.day}日"


def format_short_cn_range(start: date, end: date, workdays: int) -> str:
    return f"{start.month}月{start.day}日-{end.month}月{end.day}日（{workdays}天）"


def natural_shot_sort_key(shot_no: str) -> tuple[int, int]:
    match = re.match(r"(\d+)-(\d+)", shot_no)
    if not match:
        return (10_000, 10_000)
    return (int(match.group(1)), int(match.group(2)))


def scene_sort_key(scene_name: str) -> int:
    match = re.search(r"第(\d+)场", scene_name)
    return int(match.group(1)) if match else 10_000


def normalize_week_key(label: str) -> str:
    match = re.search(r"第(\d+)周", label)
    return f"W{match.group(1)}" if match else ""


def percent(done: int, total: int) -> int:
    return round((done / total) * 100) if total else 0


def format_scene_summary(scene_counter: Counter[str], limit: int = 3) -> str:
    ordered = sorted(scene_counter.items(), key=lambda item: (-item[1], scene_sort_key(item[0])))
    parts = [f"{name}{count}镜" for name, count in ordered[:limit]]
    if len(ordered) > limit:
        remaining = sum(count for _, count in ordered[limit:])
        if remaining:
            parts.append(f"其余{remaining}镜")
    return " + ".join(parts)


def format_type_summary(type_counter: Counter[str]) -> str:
    ordered = [
        ("keyframe", "关键帧"),
        ("scene", "场景图"),
        ("normal", "普通图"),
    ]
    parts = [f"{label}{type_counter.get(key, 0)}镜" for key, label in ordered if type_counter.get(key, 0)]
    return " / ".join(parts) if parts else "无排期"


def choose_scene_note(remaining: int, schedule_days: int, type_counter: Counter[str]) -> str:
    if remaining == 0:
        return "已从执行排期中移除"
    dominant = max(type_counter.items(), key=lambda item: item[1])[0] if type_counter else "normal"
    dominant_label = {
        "keyframe": "关键帧优先",
        "scene": "场景图穿插",
        "normal": "普通图收尾",
    }[dominant]
    return f"剩 {remaining} 镜 · {schedule_days} 个排期日 · {dominant_label}"


def parse_storyboard_workbook(path: Path) -> tuple[list[dict], list[str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.worksheets[0]
    header = [clean(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]

    rows = []
    scene_order = []
    seen_scenes = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = [clean(value) for value in row[: len(header)]]
        if not values[0] or not values[1]:
            continue

        record = dict(zip(header, values))
        scene_name = record["分场"]
        if scene_name not in seen_scenes:
            seen_scenes.add(scene_name)
            scene_order.append(scene_name)

        raw_status = record.get("状态", "")
        generated_image = record.get("讨论", "") or record.get("后期处理画面", "")

        rows.append(
            {
                "sceneName": scene_name,
                "shotNo": record["镜头号"],
                "shotId": record["镜头号"],
                "description": record.get("文字描述（画面描述，景别，摄影机角度）", ""),
                "dialogue": record.get("台词", ""),
                "referenceImage": record.get("参考画面", ""),
                "finalImage": generated_image,
                "statusLabel": raw_status or "未开始",
                "status": STATUS_LABEL_TO_KEY.get(raw_status, "pending"),
                "owner": record.get("制作者", "") or "未分配",
                "direction": record.get("方向", ""),
                "directorNote": record.get("导演修改意见", ""),
                "postImage": record.get("后期处理画面", ""),
                "mixNote": record.get("MIX混声细节", ""),
                "cameraMove": record.get("运镜", ""),
                "performanceNote": record.get("补充视频表演状态", ""),
                "pendingVideo": record.get("待检验视频素材 命名规则为场号-镜号", ""),
                "videoNote": record.get("视频修改意见", ""),
                "finalVideo": record.get("最终确定视频素材 U-场号-镜号", ""),
                "videoPrompt": record.get("视频提示词", ""),
                "reviewNote": record.get("导演与制片核验", ""),
                "parentRecord": record.get("父记录", ""),
                "isKeyframe": bool(record.get("关键帧", "")),
                "extraTag": record.get("其他", ""),
            }
        )

    return rows, scene_order


def parse_schedule_workbook(path: Path) -> dict:
    workbook = load_workbook(path, data_only=True)
    summary_sheet = workbook.worksheets[0]

    title = clean(summary_sheet["A1"].value) or "湖村漫画转分镜 — 每日详细工作排期"
    range_line = clean(summary_sheet["A2"].value)
    range_match = re.search(
        r"排期周期：(?P<start>\d{4}/\d{1,2}/\d{1,2})\s*[—-]\s*(?P<end>\d{4}/\d{1,2}/\d{1,2}).*共(?P<days>\d+)个工作日",
        range_line,
    )
    if not range_match:
        raise ValueError("Could not parse the schedule date range from the workbook header")

    start_date = parse_yyyy_mm_dd(range_match.group("start"))
    end_date = parse_yyyy_mm_dd(range_match.group("end"))
    total_workdays = int(range_match.group("days"))

    items = []
    member_names = []
    for sheet in workbook.worksheets[1:]:
        member = clean(sheet.title).split("-")[0]
        member_names.append(member)

        current_date = None
        current_week_label = ""
        current_week_key = ""
        current_weekday = ""

        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = [clean(value) for value in row[:9]]
            if not any(values):
                continue

            first = values[0]
            if first.startswith("══════"):
                match = re.search(r"第\d+周", first)
                if match:
                    current_week_label = match.group(0)
                    current_week_key = normalize_week_key(current_week_label)
                continue

            if first.startswith("📊"):
                continue

            raw_date, raw_weekday, raw_week, _, scene_name, shot_no, type_label, status_label, note = values
            if raw_date:
                current_date = parse_mmdd(raw_date)
            if raw_weekday:
                current_weekday = raw_weekday
            if raw_week:
                current_week_label = raw_week
                current_week_key = normalize_week_key(raw_week)

            if not scene_name or not shot_no or current_date is None:
                continue

            type_info = TYPE_META.get(type_label, {"key": "normal", "short": type_label or "普通图"})
            normalized_status = STATUS_LABEL_TO_KEY.get(status_label, "pending")

            items.append(
                {
                    "date": iso_date(current_date),
                    "dateLabel": current_date.strftime("%m/%d"),
                    "weekday": current_weekday or chinese_weekday(current_date),
                    "week": current_week_key,
                    "weekLabel": current_week_label,
                    "member": member,
                    "sceneName": scene_name,
                    "shotNo": shot_no,
                    "typeLabel": type_label,
                    "typeKey": type_info["key"],
                    "typeShort": type_info["short"],
                    "statusLabel": status_label or "未开始",
                    "status": normalized_status,
                    "note": note,
                }
            )

    items.sort(key=lambda item: (item["date"], item["member"], natural_shot_sort_key(item["shotNo"])))

    daily_groups = defaultdict(list)
    weekly_groups = defaultdict(list)
    scene_groups = defaultdict(list)
    for item in items:
        daily_groups[item["date"]].append(item)
        weekly_groups[item["week"]].append(item)
        scene_groups[item["sceneName"]].append(item)

    daily_schedule = []
    for day in sorted(daily_groups):
        day_items = daily_groups[day]
        member_rows = []
        for member in member_names:
            member_items = [item for item in day_items if item["member"] == member]
            if not member_items:
                continue
            scene_counter = Counter(item["sceneName"] for item in member_items)
            type_counter = Counter(item["typeKey"] for item in member_items)
            member_rows.append(
                {
                    "member": member,
                    "shots": len(member_items),
                    "statusCounts": dict(Counter(item["status"] for item in member_items)),
                    "typeCounts": dict(type_counter),
                    "sceneSummary": format_scene_summary(scene_counter, limit=2),
                    "items": [
                        {
                            "sceneName": item["sceneName"],
                            "shotNo": item["shotNo"],
                            "typeLabel": item["typeLabel"],
                            "statusLabel": item["statusLabel"],
                            "note": item["note"],
                        }
                        for item in member_items
                    ],
                }
            )

        day_type_counter = Counter(item["typeKey"] for item in day_items)
        day_scene_counter = Counter(item["sceneName"] for item in day_items)
        day_status_counter = Counter(item["status"] for item in day_items)
        focus_note = "关键帧优先" if day_type_counter.get("keyframe", 0) else "普通图推进"
        if day_type_counter.get("scene", 0):
            focus_note += " · 含场景图"

        daily_schedule.append(
            {
                "date": day,
                "label": day_items[0]["dateLabel"],
                "weekday": day_items[0]["weekday"],
                "week": day_items[0]["week"],
                "weekLabel": day_items[0]["weekLabel"],
                "totalShots": len(day_items),
                "sceneSummary": format_scene_summary(day_scene_counter, limit=3),
                "typeSummary": format_type_summary(day_type_counter),
                "focusNote": focus_note,
                "statusCounts": dict(day_status_counter),
                "members": member_rows,
            }
        )

    return {
        "title": title,
        "startDate": start_date,
        "endDate": end_date,
        "workdays": total_workdays,
        "items": items,
        "dailySchedule": daily_schedule,
        "weeklyGroups": weekly_groups,
        "sceneGroups": scene_groups,
        "memberNames": member_names,
        "scheduleShotMap": {item["shotNo"]: item for item in items},
    }


def merge_storyboard_and_schedule(storyboard_rows: list[dict], scene_order: list[str], schedule_data: dict) -> dict:
    schedule_shot_map = schedule_data["scheduleShotMap"]
    all_rows = []
    schedule_scene_groups = schedule_data["sceneGroups"]

    for row in storyboard_rows:
        shot_no = row["shotNo"]
        schedule_item = schedule_shot_map.get(shot_no)

        if row["owner"] == "LU":
            final_status = "done"
            final_status_label = "完成"
            final_owner = "LU"
        elif schedule_item:
            final_status = schedule_item["status"]
            final_status_label = schedule_item["statusLabel"]
            final_owner = schedule_item["member"]
        else:
            final_status = "done"
            final_status_label = "完成"
            final_owner = row["owner"] or "未分配"

        all_rows.append(
            {
                **row,
                "owner": final_owner,
                "status": final_status,
                "statusLabel": final_status_label,
                "plannedDate": schedule_item["date"] if schedule_item else "",
                "plannedDateLabel": schedule_item["dateLabel"] if schedule_item else "",
                "plannedWeek": schedule_item["week"] if schedule_item else "",
                "plannedWeekLabel": schedule_item["weekLabel"] if schedule_item else "",
                "plannedTypeLabel": schedule_item["typeLabel"] if schedule_item else "",
                "plannedTypeKey": schedule_item["typeKey"] if schedule_item else "",
                "plannedNote": schedule_item["note"] if schedule_item else "",
            }
        )

    coverage = {
        "textCoverage": sum(1 for row in all_rows if row["description"] or row["dialogue"]),
        "referenceCoverage": sum(1 for row in all_rows if row["referenceImage"]),
        "finalImageCoverage": sum(1 for row in all_rows if row["finalImage"]),
        "finalVideoCoverage": sum(1 for row in all_rows if row["finalVideo"]),
    }

    grouped_rows = defaultdict(list)
    for row in all_rows:
        grouped_rows[row["sceneName"]].append(row)

    scenes = []
    for order, scene_name in enumerate(scene_order, start=1):
        shots = sorted(grouped_rows[scene_name], key=lambda item: natural_shot_sort_key(item["shotNo"]))
        done_shots = sum(1 for shot in shots if shot["status"] == "done")
        revision_shots = sum(1 for shot in shots if shot["status"] == "revision")
        pending_shots = sum(1 for shot in shots if shot["status"] == "pending")

        schedule_items = sorted(
            schedule_scene_groups.get(scene_name, []),
            key=lambda item: (item["date"], item["member"], natural_shot_sort_key(item["shotNo"])),
        )
        week_keys = []
        for item in schedule_items:
            if item["week"] and item["week"] not in week_keys:
                week_keys.append(item["week"])
        member_labels = sorted({item["member"] for item in schedule_items})
        type_counter = Counter(item["typeKey"] for item in schedule_items)
        remaining = revision_shots + pending_shots

        if remaining:
            owner_label = " / ".join(member_labels) if member_labels else "待分配"
            week_label = " · ".join(week_keys) if week_keys else "待排期"
        else:
            owner_label = "LU" if all(shot["owner"] == "LU" for shot in shots) else "已完成"
            week_label = "已完成"

        if type_counter.get("keyframe", 0):
            type_label = "关键帧优先"
        elif type_counter.get("scene", 0):
            type_label = "场景图"
        elif remaining:
            type_label = "普通图"
        else:
            type_label = "收尾完成"

        schedule_days = len({item["date"] for item in schedule_items})
        scenes.append(
            {
                "id": f"scene-{order:02d}",
                "order": order,
                "name": scene_name,
                "ownerLabel": owner_label,
                "weekLabel": week_label,
                "typeLabel": type_label,
                "note": choose_scene_note(remaining, schedule_days, type_counter),
                "totalShots": len(shots),
                "doneShots": done_shots,
                "revisionShots": revision_shots,
                "pendingShots": pending_shots,
                "progress": percent(done_shots, len(shots)),
                "shots": shots,
            }
        )

    return {
        "shots": all_rows,
        "scenes": scenes,
        "coverage": coverage,
    }


def build_week_plan(schedule_data: dict, storyboard_data: dict) -> dict:
    daily_schedule = schedule_data["dailySchedule"]
    member_names = [member for member in schedule_data["memberNames"] if member != "LU"]
    initial_done = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "done")
    total_shots = len(storyboard_data["shots"])

    week_order = sorted(schedule_data["weeklyGroups"], key=lambda key: int(key[1:]))
    cumulative = initial_done
    weeks = []
    weekly_totals = [
        {
            "week": "已完成",
            "workdays": "—",
            "targetShots": "—",
            "ownerMin": "—",
            "ownerMing": "—",
            "cumulative": f"{initial_done}镜 ({percent(initial_done, total_shots)}%)",
        }
    ]
    checkpoints = []

    for week_key in week_order:
        week_items = schedule_data["weeklyGroups"][week_key]
        week_dates = sorted({item["date"] for item in week_items})
        date_objects = [date.fromisoformat(value) for value in week_dates]
        start = min(date_objects)
        end = max(date_objects)
        week_total = len(week_items)
        cumulative += week_total

        week_scene_counter = Counter(item["sceneName"] for item in week_items)
        week_type_counter = Counter(item["typeKey"] for item in week_items)

        goal_prefix = "关键帧优先"
        if week_type_counter.get("scene", 0) > week_type_counter.get("keyframe", 0):
            goal_prefix = "场景图穿插推进"
        if week_type_counter.get("normal", 0) > week_type_counter.get("keyframe", 0) + week_type_counter.get("scene", 0):
            goal_prefix = "普通图收尾冲刺"

        assignments = []
        assignment_lookup = {}
        for member in member_names:
            member_items = [item for item in week_items if item["member"] == member]
            if not member_items:
                continue
            member_scene_counter = Counter(item["sceneName"] for item in member_items)
            member_type_counter = Counter(item["typeKey"] for item in member_items)
            member_dates = defaultdict(list)
            for item in member_items:
                member_dates[item["date"]].append(item)

            notes = []
            for day in sorted(member_dates):
                day_items = member_dates[day]
                scene_counter = Counter(item["sceneName"] for item in day_items)
                type_counter = Counter(item["typeKey"] for item in day_items)
                note_suffix = " · 先清修改中" if any(item["status"] == "revision" for item in day_items) else ""
                notes.append(
                    f"{day_items[0]['dateLabel']}：{format_scene_summary(scene_counter, limit=2)} · {format_type_summary(type_counter)}{note_suffix}"
                )

            shots = len(member_items)
            daily_average = shots / len(member_dates)
            assignment = {
                "member": member,
                "task": f"{format_scene_summary(member_scene_counter, limit=3)}（{format_type_summary(member_type_counter)}）",
                "shots": str(shots),
                "dailyAverage": f"{daily_average:.1f}",
                "successMarker": f"共 {shots} 镜 · 日均 {daily_average:.1f}",
                "notes": notes,
            }
            assignments.append(assignment)
            assignment_lookup[member] = assignment

        heavier_member = max(assignments, key=lambda item: int(item["shots"])) if assignments else None
        lighter_member = min(assignments, key=lambda item: int(item["shots"])) if len(assignments) > 1 else None
        gap_note = ""
        if heavier_member and lighter_member and heavier_member is not lighter_member:
            gap = int(heavier_member["shots"]) - int(lighter_member["shots"])
            gap_note = f"{heavier_member['member']} 本周多扛 {gap} 镜，周中若掉速，优先调换普通图。"

        week = {
            "week": week_key,
            "dateRange": format_short_cn_range(start, end, len(week_dates)),
            "goal": f"{goal_prefix} · {week_total}镜 · {format_scene_summary(week_scene_counter, limit=2)}",
            "nature": f"{format_type_summary(week_type_counter)}，本周主压 {format_scene_summary(week_scene_counter, limit=2)}。",
            "notes": [gap_note] if gap_note else [],
            "weekend": "",
            "targetShots": f"{week_total}镜",
            "workdays": f"{len(week_dates)}天",
            "cumulative": f"{cumulative}镜 ({percent(cumulative, total_shots)}%)",
            "assignments": assignments,
        }
        weeks.append(week)

        weekly_totals.append(
            {
                "week": week_key,
                "workdays": week["workdays"],
                "targetShots": week["targetShots"],
                "ownerMin": f"{assignment_lookup.get('敏', {}).get('shots', '0')}镜",
                "ownerMing": f"{assignment_lookup.get('铭', {}).get('shots', '0')}镜",
                "cumulative": week["cumulative"],
            }
        )

        heaviest_scene = format_scene_summary(week_scene_counter, limit=1)
        checkpoints.append(
            {
                "date": f"{end.month}月{end.day}日（{chinese_weekday(end)}）",
                "target": f"累计 {cumulative}/{total_shots} · {week_total}镜",
                "riskSignal": f"若 {heaviest_scene} 未按时推进，下周排期会被直接压缩。",
                "response": gap_note or "优先把关键帧与修改中的镜头做完，再补普通图。",
                "weekKey": week_key,
            }
        )

    remaining_scene_counter = Counter()
    for shot in storyboard_data["shots"]:
        if shot["status"] != "done" and shot["owner"] != "LU":
            remaining_scene_counter[shot["sceneName"]] += 1

    heaviest_remaining = format_scene_summary(remaining_scene_counter, limit=3)
    plan_b = [
        "若当周落后超过 8 镜，先保关键帧和修改中的镜头，普通图顺延。",
        f"后半程优先守住 {heaviest_remaining or '大场收尾'}，避免最后一周集中爆量。",
        "敏 / 铭 的普通图可互相调换，但关键帧尽量不跨人反复切换。",
    ]

    return {
        "weeks": weeks,
        "checkpoints": checkpoints,
        "weeklyTotals": weekly_totals,
        "dailySchedule": daily_schedule,
        "routines": [],
        "planB": plan_b,
    }


def build_owner_load(storyboard_data: dict, week_plan: dict) -> list[dict]:
    planned_by_member = defaultdict(list)
    for week in week_plan["weeks"]:
        for assignment in week["assignments"]:
            planned_by_member[assignment["member"]].append(
                {
                    "week": week["week"],
                    "task": assignment["task"],
                    "shots": assignment["shots"],
                    "dailyAverage": assignment["dailyAverage"],
                    "successMarker": assignment["successMarker"],
                    "notes": assignment.get("notes", []),
                }
            )

    member_groups = defaultdict(list)
    for shot in storyboard_data["shots"]:
        member_groups[shot["owner"]].append(shot)

    owner_rows = []
    for member, shots in member_groups.items():
        done_shots = sum(1 for shot in shots if shot["status"] == "done")
        revision_shots = sum(1 for shot in shots if shot["status"] == "revision")
        pending_shots = sum(1 for shot in shots if shot["status"] == "pending")
        owner_rows.append(
            {
                "member": member,
                "sceneCount": len({shot["sceneName"] for shot in shots}),
                "totalShots": len(shots),
                "doneShots": done_shots,
                "revisionShots": revision_shots,
                "pendingShots": pending_shots,
                "progress": percent(done_shots, len(shots)),
                "plannedAssignments": planned_by_member.get(member, []),
            }
        )

    owner_rows.sort(key=lambda item: (item["member"] == "LU", -item["totalShots"]))
    return owner_rows


def build_modules(storyboard_data: dict) -> list[dict]:
    total = len(storyboard_data["shots"])
    done = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "done")
    revision = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "revision")
    pending = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "pending")
    coverage = storyboard_data["coverage"]

    modules = []
    for key, config in MODULE_DEFAULTS.items():
        item = deepcopy(config)
        if key == "scriptDev":
            progress = percent(coverage["textCoverage"], total)
            metric_label = f"{coverage['textCoverage']}/{total}（已有描述）"
            status = "active"
        elif key == "assetDev":
            progress = percent(coverage["referenceCoverage"], total)
            metric_label = f"{coverage['referenceCoverage']}/{total}（参考齐）"
            status = "done"
        elif key == "storyboardDev":
            progress = percent(done, total)
            metric_label = f"{done} 完成 / {revision} 修改中 / {pending} 未开始"
            status = "active"
        else:
            progress = 0
            metric_label = f"最终视频 {coverage['finalVideoCoverage']} / 音乐未启动"
            status = "pending"

        modules.append(
            {
                "key": key,
                "title": item["title"],
                "description": item["description"],
                "progress": progress,
                "metricLabel": metric_label,
                "status": status,
                "subitems": item["subitems"],
            }
        )
    return modules


def build_project(storyboard_data: dict, week_plan: dict) -> dict:
    total = len(storyboard_data["shots"])
    done = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "done")
    revision = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "revision")
    pending = sum(1 for shot in storyboard_data["shots"] if shot["status"] == "pending")
    remaining = revision + pending
    workdays = len(week_plan["dailySchedule"])
    daily_required = round(remaining / workdays, 1) if workdays else 0

    owner_remaining = {}
    for member in ["敏", "铭"]:
        member_shots = [shot for shot in storyboard_data["shots"] if shot["owner"] == member]
        owner_remaining[member] = sum(1 for shot in member_shots if shot["status"] != "done")

    heavy_member = max(owner_remaining, key=owner_remaining.get)
    light_member = min(owner_remaining, key=owner_remaining.get)
    gap = owner_remaining[heavy_member] - owner_remaining[light_member]
    heaviest_scenes = Counter(
        shot["sceneName"] for shot in storyboard_data["shots"] if shot["status"] != "done" and shot["owner"] != "LU"
    )

    start = date.fromisoformat(week_plan["dailySchedule"][0]["date"])
    end = date.fromisoformat(week_plan["dailySchedule"][-1]["date"])

    return {
        "id": "hucun-storyboard-schedule",
        "name": "湖村",
        "title": "湖村制片工作台",
        "subtitle": "漫画转分镜每日排期总控台",
        "deadline": format_cn_date(end),
        "planDate": format_cn_date(start, include_weekday=False),
        "workdaysLabel": f"{workdays}个工作日",
        "producer": "LU",
        "makers": ["敏", "铭"],
        "episodeLabel": "3.17-4.20",
        "sceneCount": len(storyboard_data["scenes"]),
        "totalShots": total,
        "teamMembers": ["LU", "敏", "铭"],
        "summary": {
            "completed": done,
            "revision": revision,
            "pending": pending,
            "remaining": remaining,
            "completionRate": round((done / total) * 100, 1) if total else 0,
            "needsDelivery": f"{remaining}镜",
        },
        "diagnosis": {
            "summary": {
                "总镜头": str(total),
                "已完成": str(done),
                "修改中": str(revision),
                "未开始": str(pending),
                "完成率": f"{percent(done, total)}%",
                "待交付": f"{remaining}镜",
            },
            "timeAccount": f"时间账：{workdays}个工作日，执行池还剩{remaining}镜，日均需完成{daily_required}镜。",
            "rebalance": (
                f"人员负载：{heavy_member}剩 {owner_remaining[heavy_member]} 镜，"
                f"{light_member}剩 {owner_remaining[light_member]} 镜，相差 {gap} 镜。"
            ),
            "modificationNote": f"修改中 {revision} 镜已被提前排进前两周，适合先清。",
            "newWorkNote": f"未开始 {pending} 镜主要压在 {format_scene_summary(heaviest_scenes, limit=3)}。",
        },
    }


def build_payload() -> dict:
    storyboard_path = find_latest_file(STORYBOARD_NAME_HINT)
    schedule_path = find_latest_file(SCHEDULE_NAME_HINT)

    storyboard_rows, scene_order = parse_storyboard_workbook(storyboard_path)
    schedule_data = parse_schedule_workbook(schedule_path)
    storyboard_data = merge_storyboard_and_schedule(storyboard_rows, scene_order, schedule_data)
    week_plan = build_week_plan(schedule_data, storyboard_data)
    storyboard_data["ownerLoad"] = build_owner_load(storyboard_data, week_plan)
    project = build_project(storyboard_data, week_plan)
    modules = build_modules(storyboard_data)

    return {
        "project": project,
        "modules": modules,
        "storyboard": storyboard_data,
        "plan": week_plan,
        "defaults": DEFAULTS,
    }


def main() -> None:
    payload = build_payload()
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUTPUT_PATH} using {STORYBOARD_NAME_HINT} + {SCHEDULE_NAME_HINT}")


if __name__ == "__main__":
    main()
